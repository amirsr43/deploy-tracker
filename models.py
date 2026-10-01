import csv
from pathlib import Path
from database import get_connection
from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter


def create_project(name, repository="", description=""):
    connection = get_connection()

    cursor = connection.cursor()

    cursor.execute("""
        INSERT INTO projects (name, repository, description)
        VALUES (?, ?, ?)
    """, (name, repository, description))

    connection.commit()

    project_id = cursor.lastrowid

    connection.close()

    return project_id


def get_projects():
    connection = get_connection()

    cursor = connection.cursor()

    cursor.execute("""
        SELECT *
        FROM projects
        ORDER BY created_at DESC
    """)

    projects = cursor.fetchall()

    connection.close()

    return projects


def get_project(project_id):
    connection = get_connection()

    cursor = connection.cursor()

    cursor.execute("""
        SELECT *
        FROM projects
        WHERE id = ?
    """, (project_id,))

    project = cursor.fetchone()

    connection.close()

    return project


def update_project(
    project_id,
    name,
    repository="",
    description=""
):
    connection = get_connection()

    cursor = connection.cursor()

    cursor.execute("""
        UPDATE projects
        SET
            name = ?,
            repository = ?,
            description = ?
        WHERE id = ?
    """, (
        name,
        repository,
        description,
        project_id
    ))

    connection.commit()
    connection.close()


def delete_project(project_id):
    connection = get_connection()

    cursor = connection.cursor()

    cursor.execute("""
        DELETE FROM projects
        WHERE id = ?
    """, (project_id,))

    connection.commit()
    connection.close()

def create_deployment(
    project_id,
    version,
    environment,
    branch="",
    commit_hash="",
    status="Pending",
    deployed_by="",
    duration=0,
    notes="",
    error_reason=""
):
    connection = get_connection()

    cursor = connection.cursor()

    cursor.execute("""
        INSERT INTO deployments (
            project_id,
            version,
            environment,
            branch,
            commit_hash,
            status,
            deployed_by,
            duration,
            notes,
            error_reason
        )
        VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
    """, (
        project_id,
        version,
        environment,
        branch,
        commit_hash,
        status,
        deployed_by,
        duration,
        notes,
        error_reason
    ))

    connection.commit()

    deployment_id = cursor.lastrowid

    connection.close()

    return deployment_id

def get_deployments():
    connection = get_connection()

    cursor = connection.cursor()

    cursor.execute("""
        SELECT
            deployments.*,
            projects.name AS project_name
        FROM deployments
        JOIN projects
            ON projects.id = deployments.project_id
        ORDER BY deployments.deployed_at DESC
    """)

    deployments = cursor.fetchall()

    connection.close()

    return deployments

def get_deployment(deployment_id):
    connection = get_connection()

    cursor = connection.cursor()

    cursor.execute("""
        SELECT
            deployments.*,
            projects.name AS project_name
        FROM deployments
        JOIN projects
            ON projects.id = deployments.project_id
        WHERE deployments.id = ?
    """, (deployment_id,))

    deployment = cursor.fetchone()

    connection.close()

    return deployment

def update_deployment(
    deployment_id,
    project_id,
    version,
    environment,
    branch="",
    commit_hash="",
    status="Pending",
    deployed_by="",
    duration=0,
    notes="",
    error_reason=""
):
    connection = get_connection()

    cursor = connection.cursor()

    cursor.execute("""
        UPDATE deployments
        SET
            project_id = ?,
            version = ?,
            environment = ?,
            branch = ?,
            commit_hash = ?,
            status = ?,
            deployed_by = ?,
            duration = ?,
            notes = ?,
            error_reason = ?
        WHERE id = ?
    """, (
        project_id,
        version,
        environment,
        branch,
        commit_hash,
        status,
        deployed_by,
        duration,
        notes,
        error_reason,
        deployment_id
    ))

    connection.commit()
    connection.close()

def delete_deployment(deployment_id):
    connection = get_connection()

    cursor = connection.cursor()

    cursor.execute("""
        DELETE FROM deployments
        WHERE id = ?
    """, (deployment_id,))

    connection.commit()
    connection.close()

def count_projects():
    connection = get_connection()
    cursor = connection.cursor()

    cursor.execute("""
        SELECT COUNT(*) AS total
        FROM projects
    """)

    result = cursor.fetchone()
    connection.close()

    return result["total"]


def count_deployments():
    connection = get_connection()
    cursor = connection.cursor()

    cursor.execute("""
        SELECT COUNT(*) AS total
        FROM deployments
    """)

    result = cursor.fetchone()
    connection.close()

    return result["total"]


def count_successful_deployments():
    connection = get_connection()
    cursor = connection.cursor()

    cursor.execute("""
        SELECT COUNT(*) AS total
        FROM deployments
        WHERE status = 'Success'
    """)

    result = cursor.fetchone()
    connection.close()

    return result["total"]


def get_success_rate():
    total = count_deployments()

    if total == 0:
        return 0

    successful = count_successful_deployments()

    return round((successful / total) * 100, 1)


def get_recent_deployments(limit=5):
    connection = get_connection()
    cursor = connection.cursor()

    cursor.execute("""
        SELECT
            deployments.*,
            projects.name AS project_name
        FROM deployments
        JOIN projects
            ON projects.id = deployments.project_id
        ORDER BY deployments.deployed_at DESC
        LIMIT ?
    """, (limit,))

    deployments = cursor.fetchall()
    connection.close()

    return deployments

def get_deployment_status_counts():
    connection = get_connection()
    cursor = connection.cursor()

    cursor.execute("""
        SELECT
            status,
            COUNT(*) AS total
        FROM deployments
        GROUP BY status
        ORDER BY total DESC
    """)

    results = cursor.fetchall()
    connection.close()

    return results

def search_deployments(
    search="",
    environment="All",
    status="All",
    project_id=None,
    start_date="",
    end_date=""
):
    connection = get_connection()

    cursor = connection.cursor()

    query = """
        SELECT
            deployments.*,
            projects.name AS project_name
        FROM deployments
        JOIN projects
            ON projects.id = deployments.project_id
        WHERE 1=1
    """

    params = []

    if search:
        query += """
            AND (
                projects.name LIKE ?
                OR deployments.version LIKE ?
                OR deployments.branch LIKE ?
                OR deployments.commit_hash LIKE ?
                OR deployments.deployed_by LIKE ?
            )
        """

        keywoard = f"%{search}%"

        params.extend([
            keywoard,
            keywoard,
            keywoard,
            keywoard,
            keywoard,
        ])

    if environment not in ("All", "Semua", "All environments"):
        query += """
            AND deployments.environment = ?
        """

        params.append(environment)
    
    if status not in ("All", "Semua"):
        query += """
            AND deployments.status = ?
        """

        params.append(status)

    if project_id is not None:
        query += " AND deployments.project_id = ?"
        params.append(project_id)

    if start_date:
        query += " AND date(deployments.deployed_at) >= date(?)"
        params.append(start_date)

    if end_date:
        query += " AND date(deployments.deployed_at) <= date(?)"
        params.append(end_date)

    query += """
        ORDER BY deployments.deployed_at DESC
    """

    cursor.execute(query, params)

    deployments = cursor.fetchall()

    connection.close()

    return deployments

def _get_export_deployments(
    search="",
    environment="All",
    status="All",
    project_id=None,
    start_date="",
    end_date=""
):
    return search_deployments(
        search=search,
        environment=environment,
        status=status,
        project_id=project_id,
        start_date=start_date,
        end_date=end_date
    )


def export_deployments_to_csv(
    file_path,
    search="",
    environment="All",
    status="All",
    project_id=None,
    start_date="",
    end_date=""
):
    deployments = _get_export_deployments(
        search,
        environment,
        status,
        project_id,
        start_date,
        end_date
    )

    with open(
        file_path,
        "w",
        newline="",
        encoding="utf-8-sig"
    ) as file:

        writer = csv.writer(file)

        writer.writerow([
            "Project",
            "Version",
            "Environment",
            "Branch",
            "Commit Hash",
            "Status",
            "Deployed By",
            "Duration (seconds)",
            "Notes",
            "Failure Reason",
            "Deployed At"
        ])

        for deployment in deployments:
            writer.writerow([
                deployment["project_name"],
                deployment["version"],
                deployment["environment"],
                deployment["branch"] or "",
                deployment["commit_hash"] or "",
                deployment["status"],
                deployment["deployed_by"] or "",
                deployment["duration"] or 0,
                deployment["notes"] or "",
                deployment["error_reason"] or "",
                deployment["deployed_at"]
            ])

def export_deployments_to_excel(
    file_path,
    search="",
    environment="All",
    status="All",
    project_id=None,
    start_date="",
    end_date=""
):
    deployments = _get_export_deployments(
        search,
        environment,
        status,
        project_id,
        start_date,
        end_date
    )

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Deployments"

    headers = [
        "Project",
        "Version",
        "Environment",
        "Branch",
        "Commit Hash",
        "Status",
        "Deployed By",
        "Duration (seconds)",
        "Notes",
        "Failure Reason",
        "Deployed_at"
    ]

    worksheet.append(headers)

    for cell in worksheet[1]:
        cell.font = Font(bold=True)
    
    for deployment in deployments:
        worksheet.append([
            deployment["project_name"],
            deployment["version"],
            deployment["environment"],
            deployment["branch"] or "",
            deployment["commit_hash"] or "",
            deployment["status"],
            deployment["deployed_by"] or "",
            deployment["duration"] or 0,
            deployment["notes"] or "",
            deployment["error_reason"] or "",
            deployment["deployed_at"]
        ])
    
    worksheet.auto_filter.ref = worksheet.dimensions

    worksheet.freeze_panes = "A2"

    for column in worksheet.columns:
        max_length = 0
        column_letter = get_column_letter(column[0].column)

        for cell in column:
            value = str(cell.value or "")
            max_length = max(
                max_length,
                len(value)
            )
        
        worksheet.column_dimensions[
            column_letter
        ].width = min(
            max_length + 2,
            40
        )
    
    workbook.save(file_path)


def get_setting(key, default=""):
    connection = get_connection()
    cursor = connection.cursor()
    cursor.execute("SELECT value FROM settings WHERE key = ?", (key,))
    result = cursor.fetchone()
    connection.close()
    return result["value"] if result else default


def set_setting(key, value):
    connection = get_connection()
    connection.execute(
        "INSERT INTO settings (key, value) VALUES (?, ?) "
        "ON CONFLICT(key) DO UPDATE SET value = excluded.value",
        (key, value)
    )
    connection.commit()
    connection.close()

def get_deployments_by_environment():
    connection = get_connection()
    cursor = connection.cursor()

    cursor.execute("""
        SELECT
            environment,
            COUNT(*) AS total
        FROM deployments
        GROUP BY environment
        ORDER BY total DESC
    """)

    results = cursor.fetchall()
    connection.close()

    return results

def get_deployments_by_status():
    connection = get_connection()
    cursor = connection.cursor()

    cursor.execute("""
        SELECT
            status,
            COUNT(*) AS total
        FROM deployments
        GROUP BY status
        ORDER BY total DESC
    """)

    results = cursor.fetchall()
    connection.close()

    return results